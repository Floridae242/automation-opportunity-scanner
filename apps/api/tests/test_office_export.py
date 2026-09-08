from io import BytesIO

from docx import Document
from pptx import Presentation

from aos_api.office_export import render_docx, render_slides


def snapshot() -> dict[str, object]:
    return {
        "process": "Invoice intake",
        "project": "Finance operations",
        "source_version_no": 3,
        "review_status": "reviewed",
        "schema_version": "aos-report-v1",
        "generated_at": "2026-09-08T10:00:00Z",
        "pain_points": [
            {
                "category": "manual_entry",
                "description": "Staff rekeys supplier invoices.",
                "severity": "high",
            }
        ],
        "opportunities": [
            {
                "title": "Validate invoice fields automatically",
                "score": 82,
                "scoring_version": "aos-score-v1",
                "priority_band": "high",
                "confidence": "high",
                "result_state": "reviewed",
                "recommendation": {
                    "patterns": ["workflow_automation"],
                    "rationale": "Rules can validate complete invoices.",
                    "human_control": "Approve exceptions.",
                },
            }
        ],
    }


def test_docx_export_contains_snapshot_content():
    document = Document(BytesIO(render_docx(snapshot())))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Invoice intake" in text
    assert "Staff rekeys supplier invoices." in text
    assert "Validate invoice fields automatically" in text


def test_docx_export_removes_xml_invalid_characters():
    report_snapshot = snapshot() | {"process": "Invoice\x00intake"}
    document = Document(BytesIO(render_docx(report_snapshot)))
    assert "Invoiceintake" in document.paragraphs[0].text


def test_slide_export_contains_snapshot_content():
    presentation = Presentation(BytesIO(render_slides(snapshot())))
    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert len(presentation.slides) == 4
    assert "Invoice intake" in slide_text
    assert "Validate invoice fields automatically" in slide_text
