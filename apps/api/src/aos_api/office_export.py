"""DOCX and PowerPoint renderers for immutable report snapshots."""

from io import BytesIO
from typing import Any

from docx import Document
from pptx import Presentation


def _text(value: object, fallback: str = "Unknown") -> str:
    if value is None:
        return fallback
    return "".join(character for character in str(value) if _is_xml_character(character))


def _is_xml_character(character: str) -> bool:
    code_point = ord(character)
    return (
        code_point in {0x9, 0xA, 0xD}
        or 0x20 <= code_point <= 0xD7FF
        or 0xE000 <= code_point <= 0xFFFD
        or 0x10000 <= code_point <= 0x10FFFF
    )


def _records(snapshot: dict[str, object], key: str) -> list[dict[str, object]]:
    raw = snapshot.get(key)
    return [item for item in raw if isinstance(item, dict)] if isinstance(raw, list) else []


def _report_title(snapshot: dict[str, object]) -> str:
    return f"{_text(snapshot.get('process'))} automation opportunity report"


def _summary_lines(snapshot: dict[str, object]) -> list[str]:
    return [
        f"Project: {_text(snapshot.get('project'))}",
        f"Source version: {_text(snapshot.get('source_version_no'))} "
        f"({_text(snapshot.get('review_status'))})",
        f"Report schema: {_text(snapshot.get('schema_version'))}",
        f"Generated: {_text(snapshot.get('generated_at'))}",
    ]


def _opportunity_lines(opportunity: dict[str, object]) -> list[str]:
    lines = [
        f"Score: {_text(opportunity.get('score'))} ({_text(opportunity.get('scoring_version'))})",
        f"Priority: {_text(opportunity.get('priority_band'))}; "
        f"confidence: {_text(opportunity.get('confidence'))}",
        f"Evidence state: {_text(opportunity.get('result_state'))}",
    ]
    recommendation = opportunity.get("recommendation")
    if isinstance(recommendation, dict):
        patterns = recommendation.get("patterns")
        if isinstance(patterns, list):
            lines.append(f"Recommended patterns: {', '.join(map(str, patterns))}")
        lines.append(f"Rationale: {_text(recommendation.get('rationale'))}")
        lines.append(f"Human control: {_text(recommendation.get('human_control'))}")
    return lines


def render_docx(snapshot: dict[str, object]) -> bytes:
    """Render a report snapshot without reading live analysis data."""
    document = Document()
    document.add_heading(_report_title(snapshot), level=0)
    document.add_heading("Executive summary", level=1)
    for line in _summary_lines(snapshot):
        document.add_paragraph(line)

    document.add_heading("Pain points", level=1)
    pain_points = _records(snapshot, "pain_points")
    if not pain_points:
        document.add_paragraph("No pain points were recorded.")
    for pain_point in pain_points:
        document.add_paragraph(
            f"[{_text(pain_point.get('category'))}] {_text(pain_point.get('description'))} "
            f"(severity: {_text(pain_point.get('severity'))})",
            style="List Bullet",
        )

    document.add_heading("Opportunities", level=1)
    opportunities = _records(snapshot, "opportunities")
    if not opportunities:
        document.add_paragraph("No opportunities were recorded.")
    for index, opportunity in enumerate(opportunities, start=1):
        document.add_heading(f"{index}. {_text(opportunity.get('title'))}", level=2)
        for line in _opportunity_lines(opportunity):
            document.add_paragraph(line, style="List Bullet")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_bullet_slide(presentation: Any, title: str, lines: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line


def render_slides(snapshot: dict[str, object]) -> bytes:
    """Render a concise presentation from a report snapshot."""
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = _report_title(snapshot)
    title_slide.placeholders[1].text = f"Generated {_text(snapshot.get('generated_at'))}"
    _add_bullet_slide(presentation, "Executive summary", _summary_lines(snapshot))

    pain_lines = [
        f"[{_text(point.get('category'))}] {_text(point.get('description'))}"
        for point in _records(snapshot, "pain_points")
    ]
    _add_bullet_slide(presentation, "Pain points", pain_lines or ["No pain points were recorded."])

    opportunities = _records(snapshot, "opportunities")
    if not opportunities:
        _add_bullet_slide(presentation, "Opportunities", ["No opportunities were recorded."])
    for index, opportunity in enumerate(opportunities, start=1):
        _add_bullet_slide(
            presentation,
            f"{index}. {_text(opportunity.get('title'))}",
            _opportunity_lines(opportunity),
        )

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
